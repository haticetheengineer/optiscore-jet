import streamlit as st
import pandas as pd
import re
import io
import math
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, HRFlowable, Image as RLImage,
                                 PageBreak)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

# ── Sayfa ayarları ──────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Optik Notlandırma",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Özel CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Mono:wght@400;500&family=Syne:wght@700;800&family=DM+Sans:wght@400;500&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
.stApp { background: #0f0f13; color: #e8e6f0; }
[data-testid="stSidebar"] { background: #16151d !important; border-right: 1px solid #2a2838; }
[data-testid="stSidebar"] .block-container { padding-top: 2rem; }
.hero { padding: 2.5rem 0 1.5rem 0; border-bottom: 1px solid #2a2838; margin-bottom: 2rem; }
.hero h1 { font-family: 'Syne', sans-serif; font-size: 2.4rem; font-weight: 800; letter-spacing: -0.03em; color: #ffffff; margin: 0; line-height: 1.1; }
.hero h1 span { color: #7c6af7; }
.hero p { color: #7a7890; font-size: 0.95rem; margin-top: 0.5rem; }
.metric-row { display: grid; grid-template-columns: repeat(5, 1fr); gap: 1rem; margin: 1.5rem 0; }
.metric-card { background: #1a1926; border: 1px solid #2a2838; border-radius: 12px; padding: 1.2rem 1.4rem; }
.metric-card .label { font-size: 0.75rem; color: #7a7890; text-transform: uppercase; letter-spacing: 0.08em; font-weight: 500; }
.metric-card .value { font-family: 'DM Mono', monospace; font-size: 2rem; font-weight: 500; color: #ffffff; margin-top: 0.2rem; line-height: 1; }
.metric-card .value.accent { color: #7c6af7; }
.metric-card .value.green  { color: #4ade80; }
.metric-card .value.amber  { color: #fbbf24; }
.metric-card .value.red    { color: #f87171; }
.metric-card .value.blue   { color: #60a5fa; }
.step-badge { display: inline-flex; align-items: center; gap: 0.5rem; background: #1a1926; border: 1px solid #2a2838; border-radius: 999px; padding: 0.3rem 0.8rem; font-size: 0.78rem; color: #7c6af7; font-family: 'DM Mono', monospace; margin-bottom: 0.5rem; }
.section-title { font-family: 'Syne', sans-serif; font-size: 1.15rem; font-weight: 700; color: #fff; margin: 2rem 0 1rem 0; padding-bottom: 0.4rem; border-bottom: 1px solid #2a2838; }
.warn-box { background: #2a1f00; border-left: 3px solid #fbbf24; padding: 0.7rem 1rem; border-radius: 0 8px 8px 0; font-size: 0.85rem; color: #fbbf24; margin: 0.3rem 0; font-family: 'DM Mono', monospace; }
.ok-box { background: #0d2218; border-left: 3px solid #4ade80; padding: 0.7rem 1rem; border-radius: 0 8px 8px 0; font-size: 0.85rem; color: #4ade80; margin: 0.3rem 0; }
.stTextInput > div > div > input { background: #1a1926 !important; border: 1px solid #2a2838 !important; color: #e8e6f0 !important; font-family: 'DM Mono', monospace !important; border-radius: 8px !important; font-size: 1rem !important; }
.stButton > button { background: #7c6af7 !important; color: white !important; border: none !important; border-radius: 8px !important; font-family: 'Syne', sans-serif !important; font-weight: 700 !important; font-size: 0.95rem !important; padding: 0.6rem 1.8rem !important; transition: all 0.2s !important; }
.stButton > button:hover { background: #6b59e8 !important; transform: translateY(-1px); }
[data-testid="stDownloadButton"] > button { background: #1a2e1a !important; border: 1px solid #4ade80 !important; color: #4ade80 !important; border-radius: 8px !important; font-family: 'Syne', sans-serif !important; font-weight: 700 !important; }
[data-testid="stDownloadButton"] > button:hover { background: #0d2218 !important; }
.stFileUploader { background: #1a1926 !important; border-radius: 10px !important; }
hr { border-color: #2a2838 !important; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# CORE FONKSİYONLAR
# ══════════════════════════════════════════════════════════════════════════════

def satir_parse_et(satir):
    satir = satir.strip()
    if not satir:
        return None, None
    blok_match = re.search(r'(C?\d{8,30})', satir, re.IGNORECASE)
    if not blok_match:
        return None, f"Sayısal blok bulunamadı → {satir[:60]}"
    ad_soyad = satir[:blok_match.start()].strip()
    if not ad_soyad:
        return None, "Ad Soyad boş"
    blok     = blok_match.group(1)
    blok_end = blok_match.end()
    cap      = blok.upper().startswith('C')
    rakamlar = blok[1:] if cap else blok
    tc = ""; ogr_no = ""; cep = ""
    if len(rakamlar) >= 19:
        tc           = rakamlar[:11]
        ogr_no_rakam = rakamlar[11:19]
        kalan_rakam  = rakamlar[19:]
        if re.match(r'^(05\d{9}|5\d{9})$', kalan_rakam):
            cep = kalan_rakam if kalan_rakam.startswith('0') else '0' + kalan_rakam
        ogr_no = ('C' if cap else '') + ogr_no_rakam
    elif len(rakamlar) == 11:
        tc = rakamlar
        kalan_sonra = satir[blok_end:].strip()
        ogr_match = re.match(r'^(C?\d{8})', kalan_sonra, re.IGNORECASE)
        if not ogr_match:
            return None, f"Öğrenci no bulunamadı → {kalan_sonra[:30]}"
        ogr_no   = ogr_match.group(1).upper()
        blok_end += ogr_match.end()
        kalan2    = satir[blok_end:].strip()
        cep_match = re.match(r'^(05\d{9}|5\d{9})', kalan2)
        if cep_match:
            c = cep_match.group(1)
            cep = c if c.startswith('0') else '0' + c
            blok_end += cep_match.end()
    elif len(rakamlar) == 8:
        ogr_no = ('C' if cap else '') + rakamlar
    else:
        return None, f"Sayısal blok uzunluğu beklenmedik ({len(rakamlar)} hane) → {blok}"
    kalan    = satir[blok_end:].strip()
    kalan    = re.sub(r'\b\d{5,}\b', '', kalan)
    cevaplar = re.sub(r'[^A-Ea-e0 ]', '', kalan).replace(' ', '0').upper()
    if not cevaplar:
        return None, "Cevaplar boş"
    return {"ad_soyad": ad_soyad, "tc": tc, "ogr_no": ogr_no, "cevaplar": cevaplar, "cep": cep}, None


def ogrenci_puanla(cevaplar, anahtar):
    puan = 100 / len(anahtar)
    cevaplar = cevaplar.ljust(len(anahtar), '0')[:len(anahtar)]
    sonuc = []
    for c, a in zip(cevaplar, anahtar):
        if a.upper() == 'X':
            sonuc.append(round(puan, 4))
        elif c not in ('0', ' ', '') and c == a:
            sonuc.append(round(puan, 4))
        else:
            sonuc.append(0)
    return sonuc


def isle(metin, anahtar):
    satirlar = metin.splitlines()
    sonuclar, hatalar = [], []
    for i, satir in enumerate(satirlar, 1):
        veri, hata = satir_parse_et(satir)
        if hata:
            hatalar.append(f"Satır {i}: {hata}")
            continue
        if veri is None:
            continue
        puanlar = ogrenci_puanla(veri["cevaplar"], anahtar)
        kayit = {"ogr_no": veri["ogr_no"], "ad_soyad": veri["ad_soyad"],
                 "tc": veri.get("tc", ""), "cep": veri.get("cep", "")}
        for s, p in enumerate(puanlar, 1):
            kayit[f"S{s}"] = p
        kayit["toplam"] = round(sum(puanlar), 2)
        kayit["cevaplar"] = veri["cevaplar"]
        sonuclar.append(kayit)
    return sonuclar, hatalar


# ══════════════════════════════════════════════════════════════════════════════
# SORU ANALİZİ
# ══════════════════════════════════════════════════════════════════════════════

def soru_analizi_hesapla(sonuclar, anahtar):
    """Her soru için istatistik hesapla."""
    n = len(sonuclar)
    soru_sayisi = len(anahtar)
    analiz = []
    for i, dogru_c in enumerate(anahtar, 1):
        col = f"S{i}"
        puan_per = 100 / soru_sayisi
        dogru_sayisi = sum(1 for r in sonuclar if r[col] > 0)
        bos_sayisi   = sum(1 for r in sonuclar
                          if len(r["cevaplar"]) < i or r["cevaplar"][i-1] in ('0', ' ', ''))
        yanlis_sayisi = n - dogru_sayisi - bos_sayisi
        # Şık dağılımı
        sik_dag = {s: 0 for s in 'ABCDE0'}
        for r in sonuclar:
            c = r["cevaplar"][i-1] if len(r["cevaplar"]) >= i else '0'
            if c in sik_dag:
                sik_dag[c] += 1
        gucluk = round(dogru_sayisi / n * 100, 1) if n > 0 else 0
        # Ayırt edicilik: üst %27 - alt %27
        toplamlar = sorted([r["toplam"] for r in sonuclar])
        k = max(1, int(n * 0.27))
        ust_sinir = toplamlar[-k] if k <= n else toplamlar[-1]
        alt_sinir = toplamlar[k-1] if k <= n else toplamlar[0]
        ust_grup = [r for r in sonuclar if r["toplam"] >= ust_sinir]
        alt_grup = [r for r in sonuclar if r["toplam"] <= alt_sinir]
        ust_dogru = sum(1 for r in ust_grup if r[col] > 0) / len(ust_grup) if ust_grup else 0
        alt_dogru = sum(1 for r in alt_grup if r[col] > 0) / len(alt_grup) if alt_grup else 0
        ayirt = round((ust_dogru - alt_dogru), 3)
        analiz.append({
            "soru": i,
            "anahtar": dogru_c,
            "dogru": dogru_sayisi,
            "yanlis": yanlis_sayisi,
            "bos": bos_sayisi,
            "gucluk": gucluk,
            "ayirt": ayirt,
            "sik_A": sik_dag.get('A', 0),
            "sik_B": sik_dag.get('B', 0),
            "sik_C": sik_dag.get('C', 0),
            "sik_D": sik_dag.get('D', 0),
            "sik_E": sik_dag.get('E', 0),
            "sik_0": sik_dag.get('0', 0),
        })
    return analiz


# ══════════════════════════════════════════════════════════════════════════════
# ŞIK EXCEL
# ══════════════════════════════════════════════════════════════════════════════

def renk(hex_str):
    return hex_str.lstrip('#')

# Renk paleti
C_BG        = "0F0F13"
C_HEADER    = "7C6AF7"
C_HEADER2   = "4C3FC7"
C_ALT       = "1A1926"
C_BORDER    = "2A2838"
C_GREEN     = "4ADE80"
C_AMBER     = "FBBF24"
C_RED       = "F87171"
C_BLUE      = "60A5FA"
C_WHITE     = "FFFFFF"
C_GRAY      = "7A7890"
C_DARK      = "16151D"

def h_fill(hex_col):
    return PatternFill("solid", fgColor=hex_col)

def h_font(hex_col, bold=False, size=10, name="Calibri"):
    return Font(color=hex_col, bold=bold, size=size, name=name)

def h_border(style="thin", hex_col=C_BORDER):
    s = Side(style=style, color=hex_col)
    return Border(left=s, right=s, top=s, bottom=s)

def h_align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def stil_baslik_satiri(ws, row, n_cols, text, fill_color=C_HEADER, font_size=13):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    c = ws.cell(row=row, column=1, value=text)
    c.fill    = h_fill(fill_color)
    c.font    = h_font(C_WHITE, bold=True, size=font_size)
    c.alignment = h_align("center")

def stil_header(ws, row, cols, fill=C_HEADER, font_col=C_WHITE, font_size=10):
    for j, val in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=val)
        c.fill      = h_fill(fill)
        c.font      = h_font(font_col, bold=True, size=font_size)
        c.alignment = h_align("center")
        c.border    = h_border()

def stil_veri(ws, row, vals, alt=False, alignments=None):
    bg = C_ALT if alt else C_BG
    for j, val in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=val)
        c.fill   = h_fill(bg)
        c.font   = h_font(C_WHITE, size=10)
        c.border = h_border()
        al = alignments[j-1] if alignments and j-1 < len(alignments) else "center"
        c.alignment = h_align(al)

def ayarla_sutun(ws, col_idx, width):
    ws.column_dimensions[get_column_letter(col_idx)].width = width

def gucluk_rengi(g):
    if g >= 70: return C_GREEN
    if g >= 40: return C_AMBER
    return C_RED

def ayirt_rengi(a):
    if a >= 0.3: return C_GREEN
    if a >= 0.1: return C_AMBER
    return C_RED


def excel_olustur(sonuclar, anahtar, analiz):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)

    soru_sayisi = len(anahtar)
    puan_per_s  = round(100 / soru_sayisi, 4)
    n           = len(sonuclar)

    # ── SHEET 1: Proliz Not Girişi ───────────────────────────────────────────
    ws1 = wb.create_sheet("Proliz Not Girişi")
    ws1.sheet_view.showGridLines = False
    ws1.freeze_panes = "B2"

    soru_baslik = [f"Soru {i} Puani" for i in range(1, soru_sayisi + 1)]
    headers = ["Ogrenci Numarasi"] + soru_baslik
    n_cols  = len(headers)

    # Başlık satırı
    stil_baslik_satiri(ws1, 1, n_cols, "PROLİZ NOT GİRİŞİ", C_HEADER, 14)
    ws1.row_dimensions[1].height = 32

    # Header
    stil_header(ws1, 2, headers, C_HEADER2)
    ws1.row_dimensions[2].height = 22

    # Veriler
    for i, r in enumerate(sonuclar):
        row = i + 3
        vals = [r["ogr_no"]] + [r[f"S{j}"] for j in range(1, soru_sayisi + 1)]
        alt  = (i % 2 == 1)
        bg   = C_ALT if alt else C_BG
        for j, val in enumerate(vals, 1):
            c = ws1.cell(row=row, column=j, value=val)
            c.fill   = h_fill(bg)
            c.border = h_border()
            c.alignment = h_align("center")
            if j == 1:
                c.font = h_font(C_BLUE, bold=True, size=10)
            else:
                # Renkli puan hücreleri
                if val > 0:
                    c.font = h_font(C_GREEN, bold=True, size=10)
                else:
                    c.font = h_font(C_RED, size=10)
        ws1.row_dimensions[row].height = 18

    ayarla_sutun(ws1, 1, 18)
    for j in range(2, n_cols + 1):
        ayarla_sutun(ws1, j, max(8, 12))

    # ── SHEET 2: Özet ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Özet")
    ws2.sheet_view.showGridLines = False
    ws2.freeze_panes = "A3"

    headers2 = ["Öğrenci Numarası", "Adı Soyadı", "TC No", "Cep Telefonu",
                "Toplam Puan", "Harf Notu", "Durum"]
    stil_baslik_satiri(ws2, 1, len(headers2), "ÖĞRENCİ SONUÇ ÖZETİ", C_HEADER, 14)
    ws2.row_dimensions[1].height = 32
    stil_header(ws2, 2, headers2, C_HEADER2)
    ws2.row_dimensions[2].height = 22

    def harf_notu(p):
        if p >= 90: return "AA"
        if p >= 80: return "BA"
        if p >= 70: return "BB"
        if p >= 60: return "CB"
        if p >= 50: return "CC"
        if p >= 45: return "DC"
        if p >= 40: return "DD"
        return "FF"

    for i, r in enumerate(sonuclar):
        row  = i + 3
        alt  = (i % 2 == 1)
        bg   = C_ALT if alt else C_BG
        harf = harf_notu(r["toplam"])
        dur  = "GEÇTİ" if r["toplam"] >= 50 else "KALDI"
        vals = [r["ogr_no"], r["ad_soyad"], r.get("tc",""), r.get("cep",""),
                r["toplam"], harf, dur]
        als  = ["center","left","center","center","center","center","center"]
        for j, val in enumerate(vals, 1):
            c = ws2.cell(row=row, column=j, value=val)
            c.fill      = h_fill(bg)
            c.border    = h_border()
            c.alignment = h_align(als[j-1])
            if j == 5:  # Toplam puan — renkli
                if r["toplam"] >= 70:   c.font = h_font(C_GREEN, bold=True)
                elif r["toplam"] >= 50: c.font = h_font(C_AMBER, bold=True)
                else:                   c.font = h_font(C_RED,   bold=True)
            elif j == 7:  # Durum
                if dur == "GEÇTİ": c.font = h_font(C_GREEN, bold=True)
                else:              c.font = h_font(C_RED,   bold=True)
            else:
                c.font = h_font(C_WHITE, size=10)
        ws2.row_dimensions[row].height = 18

    for j, w in enumerate([18, 26, 16, 16, 14, 10, 10], 1):
        ayarla_sutun(ws2, j, w)

    # ── SHEET 3: Soru Analizi ───────────────────────────────────────────────
    ws3 = wb.create_sheet("Soru Analizi")
    ws3.sheet_view.showGridLines = False
    ws3.freeze_panes = "A3"

    hdrs3 = ["Soru", "Anahtar", "Doğru", "Yanlış", "Boş",
             "Güçlük %", "Ayırt Edicilik", "A", "B", "C", "D", "E", "Boş"]
    stil_baslik_satiri(ws3, 1, len(hdrs3), "SORU BAZLI ANALİZ RAPORU", C_HEADER, 14)
    ws3.row_dimensions[1].height = 32
    stil_header(ws3, 2, hdrs3, C_HEADER2)
    ws3.row_dimensions[2].height = 22

    for i, a in enumerate(analiz):
        row = i + 3
        alt = (i % 2 == 1)
        bg  = C_ALT if alt else C_BG
        vals = [a["soru"], a["anahtar"], a["dogru"], a["yanlis"], a["bos"],
                a["gucluk"], a["ayirt"],
                a["sik_A"], a["sik_B"], a["sik_C"], a["sik_D"], a["sik_E"], a["sik_0"]]
        for j, val in enumerate(vals, 1):
            c = ws3.cell(row=row, column=j, value=val)
            c.fill      = h_fill(bg)
            c.border    = h_border()
            c.alignment = h_align("center")
            if j == 6:  # Güçlük
                fc = gucluk_rengi(val)
                c.font = h_font(fc, bold=True)
            elif j == 7:  # Ayırt edicilik
                fc = ayirt_rengi(val)
                c.font = h_font(fc, bold=True)
            elif j == 2:  # Anahtar
                c.font = h_font(C_AMBER, bold=True, size=11)
            elif j == 3:  # Doğru
                c.font = h_font(C_GREEN, bold=True)
            elif j == 4:  # Yanlış
                c.font = h_font(C_RED, bold=True)
            else:
                c.font = h_font(C_WHITE, size=10)
            # Şık kolonları — doğru şıkı vurgula
            if j in (8, 9, 10, 11, 12):
                sik = 'ABCDE'[j-8]
                if sik == a["anahtar"]:
                    c.fill = h_fill("1A3A1A")
                    c.font = h_font(C_GREEN, bold=True)
        ws3.row_dimensions[row].height = 18

    for j, w in enumerate([8, 10, 10, 10, 8, 12, 16, 8, 8, 8, 8, 8, 8], 1):
        ayarla_sutun(ws3, j, w)

    # ── SHEET 4: Detaylı Liste ──────────────────────────────────────────────
    ws4 = wb.create_sheet("Detaylı Liste")
    ws4.sheet_view.showGridLines = False
    ws4.freeze_panes = "F3"

    hdrs4 = (["Öğrenci No", "Adı Soyadı", "TC No", "Cep", "Toplam"] +
             [f"S{i}" for i in range(1, soru_sayisi+1)])
    stil_baslik_satiri(ws4, 1, len(hdrs4), "DETAYLI SONUÇ LİSTESİ", C_HEADER, 14)
    ws4.row_dimensions[1].height = 32
    stil_header(ws4, 2, hdrs4, C_HEADER2)
    ws4.row_dimensions[2].height = 22

    for i, r in enumerate(sonuclar):
        row = i + 3
        alt = (i % 2 == 1)
        bg  = C_ALT if alt else C_BG
        vals = ([r["ogr_no"], r["ad_soyad"], r.get("tc",""), r.get("cep",""), r["toplam"]] +
                [r[f"S{j}"] for j in range(1, soru_sayisi+1)])
        for j, val in enumerate(vals, 1):
            c = ws4.cell(row=row, column=j, value=val)
            c.fill      = h_fill(bg)
            c.border    = h_border()
            c.alignment = h_align("left" if j == 2 else "center")
            if j == 5:
                if r["toplam"] >= 70:   c.font = h_font(C_GREEN, bold=True)
                elif r["toplam"] >= 50: c.font = h_font(C_AMBER, bold=True)
                else:                   c.font = h_font(C_RED,   bold=True)
            elif j > 5:
                c.font = h_font(C_GREEN if val > 0 else C_RED, size=9)
            else:
                c.font = h_font(C_WHITE, size=10)
        ws4.row_dimensions[row].height = 18

    for j, w in enumerate([16, 24, 14, 14, 11], 1):
        ayarla_sutun(ws4, j, w)
    for j in range(6, len(hdrs4)+1):
        ayarla_sutun(ws4, j, 8)

    # ── SHEET 5: İşlem Özeti ───────────────────────────────────────────────
    ws5 = wb.create_sheet("İşlem Özeti")
    ws5.sheet_view.showGridLines = False

    toplamlar = [r["toplam"] for r in sonuclar]
    gecme = sum(1 for t in toplamlar if t >= 50)
    ort   = sum(toplamlar) / n
    std   = math.sqrt(sum((t-ort)**2 for t in toplamlar) / n)

    stil_baslik_satiri(ws5, 1, 4, "İŞLEM ÖZETİ & İSTATİSTİKLER", C_HEADER, 14)
    ws5.row_dimensions[1].height = 32
    ws5.merge_cells("A2:D2")
    ws5.row_dimensions[2].height = 8

    ozet_data = [
        ("Cevap Anahtarı",          anahtar,           C_AMBER, "center"),
        ("Soru Sayısı",             soru_sayisi,        C_WHITE, "center"),
        ("Soru Başına Puan",        puan_per_s,         C_WHITE, "center"),
        ("Toplam Öğrenci",          n,                  C_BLUE,  "center"),
        ("Geçen (≥50)",             gecme,              C_GREEN, "center"),
        ("Kalan (<50)",             n - gecme,          C_RED,   "center"),
        ("Geçme Oranı",             f"%{gecme/n*100:.1f}" if n else "—", C_AMBER, "center"),
        ("Sınıf Ortalaması",        f"{ort:.2f}",       C_GREEN, "center"),
        ("Standart Sapma",          f"{std:.2f}",       C_WHITE, "center"),
        ("En Yüksek Puan",          f"{max(toplamlar):.2f}", C_GREEN, "center"),
        ("En Düşük Puan",           f"{min(toplamlar):.2f}", C_RED,  "center"),
        ("Medyan",                  f"{sorted(toplamlar)[n//2]:.2f}", C_WHITE, "center"),
    ]

    stil_header(ws5, 3, ["Bilgi", "Değer", "", ""], C_HEADER2)
    for i, (bilgi, deger, fc, al) in enumerate(ozet_data):
        row = i + 4
        alt = (i % 2 == 1)
        bg  = C_ALT if alt else C_BG
        ws5.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        for j, val in enumerate([bilgi, deger], 1):
            c = ws5.cell(row=row, column=j, value=val)
            c.fill   = h_fill(bg)
            c.border = h_border()
            c.alignment = h_align(al)
            c.font = h_font(fc if j == 2 else C_GRAY, bold=(j==2), size=10)
        ws5.row_dimensions[row].height = 20

    for j, w in enumerate([26, 20, 10, 10], 1):
        ayarla_sutun(ws5, j, w)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# PDF RAPORU
# ══════════════════════════════════════════════════════════════════════════════

def pdf_raporu_olustur(sonuclar, anahtar, analiz, dosya_adi=""):
    """Soru analizi PDF raporu oluştur — ReportLab ile."""
    buf = io.BytesIO()

    # Renkler
    MOR     = colors.HexColor("#7C6AF7")
    KOYU    = colors.HexColor("#0F0F13")
    LACIVERT= colors.HexColor("#16151D")
    YEŞİL   = colors.HexColor("#4ADE80")
    SARI    = colors.HexColor("#FBBF24")
    KIRMIZI = colors.HexColor("#F87171")
    MAVİ    = colors.HexColor("#60A5FA")
    GECE    = colors.HexColor("#1A1926")
    BEYAZ   = colors.white
    GRİ     = colors.HexColor("#7A7890")
    BORDER  = colors.HexColor("#2A2838")

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.8*cm, rightMargin=1.8*cm,
        topMargin=1.8*cm, bottomMargin=1.8*cm,
        title="Soru Analizi Raporu"
    )

    styles = getSampleStyleSheet()
    def ps(name, **kw):
        return ParagraphStyle(name, **kw)

    s_title  = ps("T",  fontName="Helvetica-Bold", fontSize=22, textColor=BEYAZ,
                  alignment=TA_CENTER, spaceAfter=4)
    s_sub    = ps("S",  fontName="Helvetica", fontSize=10, textColor=GRİ,
                  alignment=TA_CENTER, spaceAfter=16)
    s_h2     = ps("H2", fontName="Helvetica-Bold", fontSize=13, textColor=MOR,
                  spaceBefore=16, spaceAfter=6)
    s_body   = ps("B",  fontName="Helvetica", fontSize=9, textColor=BEYAZ,
                  leading=14)
    s_mono   = ps("M",  fontName="Courier", fontSize=9, textColor=SARI)
    s_center = ps("C",  fontName="Helvetica", fontSize=9, textColor=BEYAZ,
                  alignment=TA_CENTER)
    s_footer = ps("F",  fontName="Helvetica", fontSize=7, textColor=GRİ,
                  alignment=TA_CENTER)

    n         = len(sonuclar)
    toplamlar = [r["toplam"] for r in sonuclar]
    ort       = sum(toplamlar)/n
    std       = math.sqrt(sum((t-ort)**2 for t in toplamlar)/n)
    gecme     = sum(1 for t in toplamlar if t >= 50)
    soru_say  = len(anahtar)
    puan_per  = round(100/soru_say, 4)

    story = []

    # ── Kapak alanı ─────────────────────────────────────────────────────────
    from reportlab.platypus import KeepTogether

    # Üst başlık kutusu
    hdr_tbl = Table([[Paragraph("OPTİK NOTLANDIRMA", s_title)],
                     [Paragraph("SORU ANALİZİ RAPORU", s_title)],
                     [Paragraph(f"Cevap Anahtarı: {anahtar}", s_mono)],
                     [Paragraph(f"{soru_say} Soru  ·  {n} Öğrenci  ·  {puan_per} puan/soru", s_sub)]],
                    colWidths=[doc.width])
    hdr_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), LACIVERT),
        ("ROUNDEDCORNERS", [10]),
        ("TOPPADDING",    (0,0), (-1,-1), 8),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
        ("LEFTPADDING",   (0,0), (-1,-1), 16),
        ("RIGHTPADDING",  (0,0), (-1,-1), 16),
        ("LINEBELOW", (0,-1), (-1,-1), 2, MOR),
    ]))
    story.append(hdr_tbl)
    story.append(Spacer(1, 0.5*cm))

    # ── Özet istatistik kutusu ───────────────────────────────────────────────
    story.append(Paragraph("GENEL İSTATİSTİKLER", s_h2))
    ozet_rows = [
        [Paragraph("<b>Öğrenci Sayısı</b>", s_body),
         Paragraph(f"<b>{n}</b>", ps("x", fontName="Helvetica-Bold", fontSize=11,
                                     textColor=MAVİ, alignment=TA_CENTER)),
         Paragraph("<b>Ortalama</b>", s_body),
         Paragraph(f"<b>{ort:.2f}</b>", ps("x2", fontName="Helvetica-Bold", fontSize=11,
                                            textColor=YEŞİL, alignment=TA_CENTER))],
        [Paragraph("<b>Geçen</b>", s_body),
         Paragraph(f"<b>{gecme}</b>", ps("x3", fontName="Helvetica-Bold", fontSize=11,
                                          textColor=YEŞİL, alignment=TA_CENTER)),
         Paragraph("<b>Kalan</b>", s_body),
         Paragraph(f"<b>{n-gecme}</b>", ps("x4", fontName="Helvetica-Bold", fontSize=11,
                                            textColor=KIRMIZI, alignment=TA_CENTER))],
        [Paragraph("<b>En Yüksek</b>", s_body),
         Paragraph(f"<b>{max(toplamlar):.2f}</b>", ps("x5", fontName="Helvetica-Bold", fontSize=11,
                                                        textColor=YEŞİL, alignment=TA_CENTER)),
         Paragraph("<b>En Düşük</b>", s_body),
         Paragraph(f"<b>{min(toplamlar):.2f}</b>", ps("x6", fontName="Helvetica-Bold", fontSize=11,
                                                        textColor=KIRMIZI, alignment=TA_CENTER))],
        [Paragraph("<b>Std Sapma</b>", s_body),
         Paragraph(f"<b>{std:.2f}</b>", ps("x7", fontName="Helvetica-Bold", fontSize=11,
                                            textColor=SARI, alignment=TA_CENTER)),
         Paragraph("<b>Geçme Oranı</b>", s_body),
         Paragraph(f"<b>%{gecme/n*100:.1f}</b>", ps("x8", fontName="Helvetica-Bold", fontSize=11,
                                                       textColor=SARI, alignment=TA_CENTER))],
    ]
    w2 = doc.width / 4
    oz_tbl = Table(ozet_rows, colWidths=[w2*1.3, w2*0.7, w2*1.3, w2*0.7])
    oz_tbl.setStyle(TableStyle([
        ("BACKGROUND",  (0,0), (-1,-1), GECE),
        ("GRID",        (0,0), (-1,-1), 0.5, BORDER),
        ("TOPPADDING",  (0,0), (-1,-1), 8),
        ("BOTTOMPADDING",(0,0),(-1,-1), 8),
        ("LEFTPADDING", (0,0), (-1,-1), 10),
        ("RIGHTPADDING",(0,0), (-1,-1), 10),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [GECE, LACIVERT]),
    ]))
    story.append(oz_tbl)
    story.append(Spacer(1, 0.4*cm))

    # ── Puan dağılımı grafiği ────────────────────────────────────────────────
    story.append(Paragraph("PUAN DAĞILIMI", s_h2))
    fig, ax = plt.subplots(figsize=(7, 2.5), facecolor="#0F0F13")
    ax.set_facecolor("#1A1926")
    bins = list(range(0, 105, 10))
    counts, edges = np.histogram(toplamlar, bins=bins)
    bar_colors = []
    for e in edges[:-1]:
        if e < 50:   bar_colors.append("#F87171")
        elif e < 70: bar_colors.append("#FBBF24")
        else:        bar_colors.append("#4ADE80")
    ax.bar(edges[:-1], counts, width=9, align="edge",
           color=bar_colors, edgecolor="#2A2838", linewidth=0.5)
    ax.axvline(ort, color="#7C6AF7", linewidth=1.5, linestyle="--",
               label=f"Ort: {ort:.1f}")
    ax.axvline(50,  color="#F87171", linewidth=1,   linestyle=":",
               label="Geçme (50)")
    ax.set_xlim(0, 100)
    ax.set_xlabel("Puan", color="#7A7890", fontsize=8)
    ax.set_ylabel("Öğrenci Sayısı", color="#7A7890", fontsize=8)
    ax.tick_params(colors="#7A7890", labelsize=7)
    for sp in ax.spines.values(): sp.set_edgecolor("#2A2838")
    ax.legend(fontsize=7, facecolor="#16151D", edgecolor="#2A2838",
              labelcolor="white")
    plt.tight_layout(pad=0.3)
    img_buf = io.BytesIO()
    plt.savefig(img_buf, format="png", dpi=130, bbox_inches="tight",
                facecolor="#0F0F13")
    plt.close(fig)
    img_buf.seek(0)
    story.append(RLImage(img_buf, width=doc.width, height=6*cm))
    story.append(Spacer(1, 0.4*cm))

    # ── Soru analizi tablosu ─────────────────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("SORU BAZLI ANALİZ", s_h2))

    def gucluk_label(g):
        if g >= 70: return ("Kolay", YEŞİL)
        if g >= 40: return ("Orta",  SARI)
        return ("Zor", KIRMIZI)

    def ayirt_label(a):
        if a >= 0.3: return ("İyi",    YEŞİL)
        if a >= 0.1: return ("Orta",   SARI)
        return ("Zayıf", KIRMIZI)

    # Tablo başlığı
    ps_th = ps("TH", fontName="Helvetica-Bold", fontSize=8,
               textColor=BEYAZ, alignment=TA_CENTER)
    ps_td = ps("TD", fontName="Helvetica", fontSize=8,
               textColor=BEYAZ, alignment=TA_CENTER)
    ps_mono_s = ps("MS", fontName="Courier-Bold", fontSize=9,
                   textColor=SARI, alignment=TA_CENTER)

    ana_rows = [[
        Paragraph("Soru", ps_th), Paragraph("Ans", ps_th),
        Paragraph("Doğru", ps_th), Paragraph("Yanlış", ps_th), Paragraph("Boş", ps_th),
        Paragraph("Güçlük%", ps_th), Paragraph("Ayırt.", ps_th),
        Paragraph("A", ps_th), Paragraph("B", ps_th), Paragraph("C", ps_th),
        Paragraph("D", ps_th), Paragraph("E", ps_th),
    ]]
    row_styles = [
        ("BACKGROUND", (0,0), (-1,0), MOR),
        ("GRID",       (0,0), (-1,-1), 0.4, BORDER),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING",(0,0),(-1,-1), 4),
        ("LEFTPADDING",(0,0),(-1,-1), 4),
        ("RIGHTPADDING",(0,0),(-1,-1), 4),
    ]

    for idx, a in enumerate(analiz):
        gl, gc = gucluk_label(a["gucluk"])
        al, ac = ayirt_label(a["ayirt"])
        bg = GECE if idx % 2 == 0 else LACIVERT

        def sik_p(sik_key, sik_char):
            cnt = a[sik_key]
            fc  = "#4ADE80" if sik_char == a["anahtar"] else "#E8E6F0"
            return Paragraph(f'<font color="{fc}"><b>{cnt}</b></font>', ps_td)

        row = [
            Paragraph(f"<b>{a['soru']}</b>", ps_td),
            Paragraph(f"<b>{a['anahtar']}</b>", ps_mono_s),
            Paragraph(f'<font color="#4ADE80"><b>{a["dogru"]}</b></font>', ps_td),
            Paragraph(f'<font color="#F87171">{a["yanlis"]}</font>', ps_td),
            Paragraph(f'<font color="#7A7890">{a["bos"]}</font>',   ps_td),
            Paragraph(f'<font color="#{gucluk_rengi(a["gucluk"])}">'
                      f'<b>{a["gucluk"]}</b></font>', ps_td),
            Paragraph(f'<font color="#{ayirt_rengi(a["ayirt"])}">'
                      f'<b>{a["ayirt"]:.2f}</b></font>', ps_td),
            sik_p("sik_A", "A"), sik_p("sik_B", "B"), sik_p("sik_C", "C"),
            sik_p("sik_D", "D"), sik_p("sik_E", "E"),
        ]
        ana_rows.append(row)
        row_styles.append(("BACKGROUND", (0, idx+1), (-1, idx+1), bg))

    col_ws = [1.2*cm, 1*cm, 1.4*cm, 1.4*cm, 0.9*cm, 1.6*cm, 1.6*cm,
              0.9*cm, 0.9*cm, 0.9*cm, 0.9*cm, 0.9*cm]
    # Toplam genişliği doc.width'e sıkıştır
    total = sum(col_ws)
    if total > doc.width:
        ratio = doc.width / total
        col_ws = [w * ratio for w in col_ws]

    ana_tbl = Table(ana_rows, colWidths=col_ws, repeatRows=1)
    ana_tbl.setStyle(TableStyle(row_styles))
    story.append(ana_tbl)
    story.append(Spacer(1, 0.5*cm))

    # ── Güçlük dağılımı grafiği ──────────────────────────────────────────────
    story.append(Paragraph("SORU GÜÇLÜK & AYIRT EDİCİLİK", s_h2))
    fig2, (ax1, ax2) = plt.subplots(1, 2, figsize=(7, 2.5), facecolor="#0F0F13")
    for ax in (ax1, ax2):
        ax.set_facecolor("#1A1926")
        for sp in ax.spines.values(): sp.set_edgecolor("#2A2838")
        ax.tick_params(colors="#7A7890", labelsize=7)

    sorular = [a["soru"] for a in analiz]
    guclukler = [a["gucluk"] for a in analiz]
    ayirtlar  = [a["ayirt"]  for a in analiz]
    g_colors  = ["#4ADE80" if g >= 70 else "#FBBF24" if g >= 40 else "#F87171"
                 for g in guclukler]
    a_colors  = ["#4ADE80" if a >= 0.3 else "#FBBF24" if a >= 0.1 else "#F87171"
                 for a in ayirtlar]

    ax1.bar(sorular, guclukler, color=g_colors, edgecolor="#2A2838", linewidth=0.4)
    ax1.axhline(70, color="#4ADE80", linewidth=0.8, linestyle="--", alpha=0.6)
    ax1.axhline(40, color="#FBBF24", linewidth=0.8, linestyle="--", alpha=0.6)
    ax1.set_title("Güçlük İndeksi (%)", color="white", fontsize=8, pad=4)
    ax1.set_ylim(0, 105)
    ax1.set_xlabel("Soru", color="#7A7890", fontsize=7)

    ax2.bar(sorular, ayirtlar, color=a_colors, edgecolor="#2A2838", linewidth=0.4)
    ax2.axhline(0.3, color="#4ADE80", linewidth=0.8, linestyle="--", alpha=0.6)
    ax2.axhline(0.1, color="#FBBF24", linewidth=0.8, linestyle="--", alpha=0.6)
    ax2.set_title("Ayırt Edicilik İndeksi", color="white", fontsize=8, pad=4)
    ax2.set_xlabel("Soru", color="#7A7890", fontsize=7)

    plt.tight_layout(pad=0.5)
    img2_buf = io.BytesIO()
    plt.savefig(img2_buf, format="png", dpi=130, bbox_inches="tight",
                facecolor="#0F0F13")
    plt.close(fig2)
    img2_buf.seek(0)
    story.append(RLImage(img2_buf, width=doc.width, height=6*cm))

    # Footer
    story.append(Spacer(1, 0.5*cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER))
    story.append(Spacer(1, 0.2*cm))
    story.append(Paragraph("Optik Notlandırma Sistemi — Otomatik Üretilmiştir", s_footer))

    # Build
    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(KOYU)
        canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# ARAYÜZ
# ══════════════════════════════════════════════════════════════════════════════

st.markdown("""
<div class="hero">
  <h1>Optik <span>Notlandırma</span></h1>
  <p>TXT → Excel otomatik dönüşüm · Soru Analizi · PDF Raporu</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown('<div class="step-badge">① Cevap Anahtarı</div>', unsafe_allow_html=True)
    anahtar_input = st.text_input(
        "Cevap anahtarı",
        placeholder="örn: ABCDABCDABCD  (X=iptal)",
        label_visibility="collapsed",
    ).strip().upper()

    st.markdown('<div style="margin-top:1.5rem"></div>', unsafe_allow_html=True)
    st.markdown('<div class="step-badge">② Dosya Yükle</div>', unsafe_allow_html=True)
    yuklenen = st.file_uploader(
        "Optik okuyucu TXT dosyası",
        type=["txt"],
        label_visibility="collapsed",
    )

    st.markdown('<div style="margin-top:1.5rem"></div>', unsafe_allow_html=True)
    isle_btn = st.button("▶  Notlandır & Analiz Et", use_container_width=True)

# ── Boş durum ──────────────────────────────────────────────────────────────
if not isle_btn:
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("""
        <div style="background:#1a1926;border:1px solid #2a2838;border-radius:12px;padding:1.8rem;">
            <div style="font-family:'Syne',sans-serif;font-size:1.1rem;font-weight:700;color:#fff;margin-bottom:1rem;">📌 TXT Format</div>
            <div style="font-family:'DM Mono',monospace;font-size:0.82rem;color:#7a7890;line-height:2;">
                AD SOYAD &nbsp;|&nbsp; TC (11 hane)<br>
                Öğrenci No (8 hane veya C+8)<br>
                Cevaplar (ABCDE veya 0=boş)<br><br>
                <span style="color:#7c6af7;">Örnek:</span><br>
                ALİ YILMAZ 12345678901 20230001 ABCDABCD
            </div>
        </div>""", unsafe_allow_html=True)
    with col2:
        st.markdown("""
        <div style="background:#1a1926;border:1px solid #2a2838;border-radius:12px;padding:1.8rem;">
            <div style="font-family:'Syne',sans-serif;font-size:1.1rem;font-weight:700;color:#fff;margin-bottom:1rem;">📊 Çıktılar</div>
            <div style="font-size:0.85rem;color:#7a7890;line-height:2.2;">
                <span style="color:#7c6af7;">Excel</span> → 5 sheet: Proliz, Özet, Soru Analizi, Detay, Özet İstatistik<br>
                <span style="color:#4ade80;">PDF</span> → Soru analizi raporu + grafikler<br><br>
                <b style="color:#4ade80;">X</b> = İptal soru (herkes tam puan)
            </div>
        </div>""", unsafe_allow_html=True)

elif not anahtar_input:
    st.warning("⚠️ Lütfen sol panelden cevap anahtarını girin.")

elif not yuklenen:
    st.warning("⚠️ Lütfen sol panelden bir TXT dosyası yükleyin.")

else:
    # Dosyayı oku
    try:
        raw = yuklenen.read()
        metin = None
        for enc in ("utf-8", "cp1254", "latin-1"):
            try:
                metin = raw.decode(enc); break
            except UnicodeDecodeError:
                continue
        if metin is None:
            raise ValueError("Encoding çözümlenemedi")
    except Exception as e:
        st.error(f"Dosya okunamadı: {e}")
        st.stop()

    with st.spinner("İşleniyor…"):
        sonuclar, hatalar = isle(metin, anahtar_input)

    if not sonuclar:
        st.error("Hiç öğrenci verisi işlenemedi. Dosya formatını kontrol edin.")
        st.stop()

    analiz = soru_analizi_hesapla(sonuclar, anahtar_input)

    # ── Metrikler ────────────────────────────────────────────────────────────
    toplamlar  = [r["toplam"] for r in sonuclar]
    soru_sayisi = len(anahtar_input)
    soru_puani  = round(100 / soru_sayisi, 2)
    ort         = sum(toplamlar) / len(toplamlar)
    gecme       = sum(1 for t in toplamlar if t >= 50)
    std         = math.sqrt(sum((t-ort)**2 for t in toplamlar) / len(toplamlar))

    st.markdown(f"""
    <div class="metric-row">
        <div class="metric-card">
            <div class="label">Öğrenci</div>
            <div class="value accent">{len(sonuclar)}</div>
        </div>
        <div class="metric-card">
            <div class="label">Soru</div>
            <div class="value">{soru_sayisi} <span style="font-size:0.9rem;color:#7a7890;">× {soru_puani}pt</span></div>
        </div>
        <div class="metric-card">
            <div class="label">Ortalama</div>
            <div class="value green">{ort:.1f}</div>
        </div>
        <div class="metric-card">
            <div class="label">Geçen (≥50)</div>
            <div class="value amber">{gecme}</div>
        </div>
        <div class="metric-card">
            <div class="label">Std Sapma</div>
            <div class="value blue">{std:.1f}</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    if hatalar:
        with st.expander(f"⚠️ {len(hatalar)} satırda sorun", expanded=False):
            for h in hatalar:
                st.markdown(f'<div class="warn-box">{h}</div>', unsafe_allow_html=True)

    # ── Sekmeler ─────────────────────────────────────────────────────────────
    tab1, tab2, tab3 = st.tabs(["📋 Öğrenci Sonuçları", "📊 Soru Analizi", "📈 Grafikler"])

    with tab1:
        df_on = pd.DataFrame([
            {"Öğrenci No": r["ogr_no"], "Ad Soyad": r["ad_soyad"], "Toplam": r["toplam"]}
            for r in sonuclar
        ])
        st.dataframe(
            df_on, use_container_width=True, hide_index=True,
            column_config={
                "Toplam": st.column_config.ProgressColumn(
                    "Toplam Puan (/100)", min_value=0, max_value=100, format="%.2f"
                )
            }
        )

    with tab2:
        df_ana = pd.DataFrame([{
            "Soru": a["soru"], "Anahtar": a["anahtar"],
            "Doğru": a["dogru"], "Yanlış": a["yanlis"], "Boş": a["bos"],
            "Güçlük %": a["gucluk"], "Ayırt Edicilik": a["ayirt"],
            "A": a["sik_A"], "B": a["sik_B"], "C": a["sik_C"],
            "D": a["sik_D"], "E": a["sik_E"],
        } for a in analiz])
        st.dataframe(
            df_ana, use_container_width=True, hide_index=True,
            column_config={
                "Güçlük %": st.column_config.ProgressColumn(
                    "Güçlük %", min_value=0, max_value=100, format="%.1f%%"),
                "Ayırt Edicilik": st.column_config.NumberColumn(format="%.3f"),
            }
        )
        # Özet
        kolay = sum(1 for a in analiz if a["gucluk"] >= 70)
        orta  = sum(1 for a in analiz if 40 <= a["gucluk"] < 70)
        zor   = sum(1 for a in analiz if a["gucluk"] < 40)
        st.markdown(f"""
        <div style="display:flex;gap:1rem;margin-top:1rem;">
            <div style="background:#0d2218;border:1px solid #4ade80;border-radius:8px;padding:0.7rem 1.2rem;flex:1;text-align:center;">
                <div style="font-size:0.7rem;color:#7a7890;text-transform:uppercase;">Kolay (≥70%)</div>
                <div style="font-family:'DM Mono',monospace;font-size:1.6rem;color:#4ade80;font-weight:600;">{kolay}</div>
            </div>
            <div style="background:#2a1f00;border:1px solid #fbbf24;border-radius:8px;padding:0.7rem 1.2rem;flex:1;text-align:center;">
                <div style="font-size:0.7rem;color:#7a7890;text-transform:uppercase;">Orta (40-70%)</div>
                <div style="font-family:'DM Mono',monospace;font-size:1.6rem;color:#fbbf24;font-weight:600;">{orta}</div>
            </div>
            <div style="background:#2a0d0d;border:1px solid #f87171;border-radius:8px;padding:0.7rem 1.2rem;flex:1;text-align:center;">
                <div style="font-size:0.7rem;color:#7a7890;text-transform:uppercase;">Zor (&lt;40%)</div>
                <div style="font-family:'DM Mono',monospace;font-size:1.6rem;color:#f87171;font-weight:600;">{zor}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

    with tab3:
        col_g, col_a = st.columns(2)
        with col_g:
            fig1, ax1 = plt.subplots(figsize=(5, 3), facecolor="#0F0F13")
            ax1.set_facecolor("#1A1926")
            bins = list(range(0, 105, 10))
            counts, edges = np.histogram(toplamlar, bins=bins)
            bc = ["#F87171" if e < 50 else "#FBBF24" if e < 70 else "#4ADE80"
                  for e in edges[:-1]]
            ax1.bar(edges[:-1], counts, width=9, align="edge", color=bc,
                    edgecolor="#2A2838", linewidth=0.5)
            ax1.axvline(ort, color="#7C6AF7", linewidth=1.5, linestyle="--",
                        label=f"Ort: {ort:.1f}")
            ax1.set_title("Puan Dağılımı", color="white", fontsize=10)
            ax1.tick_params(colors="#7A7890", labelsize=8)
            for sp in ax1.spines.values(): sp.set_edgecolor("#2A2838")
            ax1.legend(fontsize=7, facecolor="#16151D", edgecolor="#2A2838",
                       labelcolor="white")
            plt.tight_layout(pad=0.3)
            st.pyplot(fig1)
            plt.close(fig1)

        with col_a:
            fig2, ax2 = plt.subplots(figsize=(5, 3), facecolor="#0F0F13")
            ax2.set_facecolor("#1A1926")
            sorular   = [a["soru"]   for a in analiz]
            guclukler = [a["gucluk"] for a in analiz]
            gc = ["#4ADE80" if g >= 70 else "#FBBF24" if g >= 40 else "#F87171"
                  for g in guclukler]
            ax2.bar(sorular, guclukler, color=gc, edgecolor="#2A2838", linewidth=0.4)
            ax2.axhline(70, color="#4ADE80", linewidth=0.8, linestyle="--", alpha=0.7)
            ax2.axhline(40, color="#FBBF24", linewidth=0.8, linestyle="--", alpha=0.7)
            ax2.set_title("Soru Güçlük İndeksi (%)", color="white", fontsize=10)
            ax2.set_xlabel("Soru No", color="#7A7890", fontsize=8)
            ax2.tick_params(colors="#7A7890", labelsize=8)
            for sp in ax2.spines.values(): sp.set_edgecolor("#2A2838")
            plt.tight_layout(pad=0.3)
            st.pyplot(fig2)
            plt.close(fig2)

        # Ayırt edicilik grafiği
        fig3, ax3 = plt.subplots(figsize=(10, 2.5), facecolor="#0F0F13")
        ax3.set_facecolor("#1A1926")
        ayirtlar = [a["ayirt"] for a in analiz]
        ac = ["#4ADE80" if a >= 0.3 else "#FBBF24" if a >= 0.1 else "#F87171"
              for a in ayirtlar]
        ax3.bar(sorular, ayirtlar, color=ac, edgecolor="#2A2838", linewidth=0.4)
        ax3.axhline(0.3, color="#4ADE80", linewidth=0.8, linestyle="--", alpha=0.7, label="İyi (≥0.3)")
        ax3.axhline(0.1, color="#FBBF24", linewidth=0.8, linestyle="--", alpha=0.7, label="Orta (≥0.1)")
        ax3.set_title("Ayırt Edicilik İndeksi (Soru Bazlı)", color="white", fontsize=10)
        ax3.set_xlabel("Soru No", color="#7A7890", fontsize=8)
        ax3.tick_params(colors="#7A7890", labelsize=8)
        for sp in ax3.spines.values(): sp.set_edgecolor("#2A2838")
        ax3.legend(fontsize=7, facecolor="#16151D", edgecolor="#2A2838", labelcolor="white")
        plt.tight_layout(pad=0.3)
        st.pyplot(fig3)
        plt.close(fig3)

    # ── İndirme butonları ─────────────────────────────────────────────────────
    st.markdown("---")
    col_dl1, col_dl2 = st.columns(2)

    with col_dl1:
        with st.spinner("Excel hazırlanıyor…"):
            excel_buf = excel_olustur(sonuclar, anahtar_input, analiz)
        st.download_button(
            label="⬇️  Excel İndir (sonuc.xlsx)",
            data=excel_buf,
            file_name="sonuc.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    with col_dl2:
        with st.spinner("PDF hazırlanıyor…"):
            pdf_buf = pdf_raporu_olustur(sonuclar, anahtar_input, analiz)
        st.download_button(
            label="📄  PDF Raporu İndir (analiz.pdf)",
            data=pdf_buf,
            file_name="analiz.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    st.markdown(
        '<div class="ok-box">✅ Excel (5 sheet) + PDF Soru Analizi raporu hazır.</div>',
        unsafe_allow_html=True
    )
